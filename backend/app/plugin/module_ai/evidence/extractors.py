"""Evidence file text extraction.

Supported formats: docx, xlsx, txt, csv.
Unsupported (returns unsupported status): pdf, png, jpg.
"""

from __future__ import annotations

from pathlib import Path


def extract_text(path: str | Path) -> tuple[str, str]:
    """Extract text content from an evidence file.

    Returns (parse_status, content):
      - "parsed", content: successfully extracted text
      - "unsupported", message: file type not supported for auto-parsing
      - "failed", empty: extraction failed
    """
    file_path = Path(path)
    suffix = file_path.suffix.lower()

    if suffix in {".txt"}:
        try:
            return "parsed", file_path.read_text(encoding="utf-8", errors="ignore")
        except Exception:
            return "failed", ""

    if suffix in {".docx"}:
        try:
            from docx import Document
            doc = Document(str(file_path))
            content = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
            if not content.strip():
                return "failed", ""
            return "parsed", content
        except Exception:
            return "failed", ""

    if suffix in {".xlsx"}:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(file_path), data_only=True)
            lines: list[str] = []
            for sheet_name in wb.sheetnames:
                ws = wb[sheet_name]
                lines.append(f"=== Sheet: {sheet_name} ===")
                for row in ws.iter_rows(values_only=True):
                    row_vals = [str(cell) if cell is not None else "" for cell in row]
                    if any(v.strip() for v in row_vals):
                        lines.append("\t".join(row_vals))
            content = "\n".join(lines)
            wb.close()
            if not content.strip():
                return "failed", ""
            return "parsed", content
        except Exception:
            return "failed", ""

    if suffix in {".csv"}:
        try:
            content = file_path.read_text(encoding="utf-8", errors="ignore")
            if not content.strip():
                return "failed", ""
            return "parsed", content
        except Exception:
            return "failed", ""

    if suffix in {".pdf", ".png", ".jpg", ".jpeg"}:
        return "unsupported", f"文件类型 {suffix} 已上传，暂不支持自动文本解析（需要 OCR/PaddleOCR 等能力）。"

    return "unsupported", f"不支持的文件类型: {suffix}"
