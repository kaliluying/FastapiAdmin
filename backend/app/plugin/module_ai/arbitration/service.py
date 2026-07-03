from __future__ import annotations

import io
import re
from datetime import date
from decimal import Decimal
from typing import Any

from app.config.setting import settings
from app.core.base_schema import AuthSchema
from app.core.exceptions import CustomException
from app.core.logger import logger
from app.plugin.module_ai.chat.rag import LangChainChatModel

from .crud import ArbitrationCaseCRUD, ArbitrationDraftCRUD
from .schema import ArbitrationDraftRequestSchema


MISSING = "【待补充】"


class ArbitrationDraftService:
    """Generate labor arbitration application drafts from structured case facts."""

    def __init__(self, auth: AuthSchema) -> None:
        self.auth = auth

    async def generate_draft(self, data: ArbitrationDraftRequestSchema) -> dict[str, Any]:
        profile = self._profile_defaults()
        case = self._merge_case_data(data, profile)
        evidence_analyses = await self._fetch_evidence_analyses(data.session_id)
        template_content = self._render_template(case=case, evidence_analyses=evidence_analyses)

        content = template_content
        if data.use_ai and self._ai_configured():
            try:
                prompt = self._build_ai_prompt(template_content)
                ai_content = await LangChainChatModel().complete(prompt)
                if ai_content and ai_content.strip():
                    content = self._sanitize_ai_content(ai_content)
            except Exception as exc:
                logger.warning(f"仲裁申请书 AI 生成失败，回退模板草稿: {exc}")

        risk_tips = [
            "本申请书为 AI 辅助草稿，提交前必须人工核对身份信息、日期、金额和证据编号。",
            "金额为初步整理结果，欠薪、加班费、经济补偿等仍需结合工资流水、考勤记录进一步核算。",
            "系统不替代律师、仲裁员或最终法律判断。",
        ]
        source_summary = {
            "session_id": data.session_id,
            "claims_count": len(case["claims"]),
            "evidence_count": len(case["evidence_items"]),
            "evidence_analysis_count": len(evidence_analyses),
            "used_ai": content != template_content,
        }
        case_id: int | None = None
        draft_id: int | None = None
        if getattr(self.auth, "db", None) is not None:
            case_obj = await self._save_case(case)
            draft_obj = await ArbitrationDraftCRUD(self.auth).create_draft(
                {
                    "case_id": case_obj.id,
                    "session_id": data.session_id,
                    "title": "劳动人事争议仲裁申请书",
                    "content": content,
                    "risk_tips": risk_tips,
                    "source_summary": source_summary,
                    "source_data": self._json_safe(case),
                    "used_ai": content != template_content,
                    "review_status": "pending",
                }
            )
            case_id = case_obj.id
            draft_id = draft_obj.id

        return {
            "case_id": case_id,
            "draft_id": draft_id,
            "title": "劳动人事争议仲裁申请书",
            "content": content,
            "risk_tips": risk_tips,
            "source_summary": source_summary,
        }

    async def list_cases(self) -> dict[str, Any]:
        cases = await ArbitrationCaseCRUD(self.auth).list_cases()
        return {
            "items": [self._case_to_dict(item) for item in cases],
            "total": len(cases),
        }

    async def list_drafts(self, *, case_id: int | None = None) -> dict[str, Any]:
        drafts = await ArbitrationDraftCRUD(self.auth).list_drafts(case_id=case_id)
        return {
            "items": [self._draft_to_dict(item) for item in drafts],
            "total": len(drafts),
        }

    async def get_draft(self, *, draft_id: int) -> dict[str, Any]:
        draft = await ArbitrationDraftCRUD(self.auth).get_draft(draft_id)
        if not draft:
            raise CustomException(msg="申请书草稿不存在")
        return {
            "case_id": draft.case_id,
            "draft_id": draft.id,
            "title": draft.title,
            "content": draft.content,
            "risk_tips": draft.risk_tips or [],
            "source_summary": draft.source_summary or {},
        }

    async def export_draft(self, *, draft_id: int, file_type: str) -> bytes:
        draft = await ArbitrationDraftCRUD(self.auth).get_draft(draft_id)
        if not draft:
            raise CustomException(msg="申请书草稿不存在")
        if file_type == "docx":
            return self._build_docx(draft.title, draft.content)
        if file_type == "pdf":
            return self._build_pdf(draft.title, draft.content)
        raise CustomException(msg="不支持的导出类型")

    async def _save_case(self, case: dict[str, Any]):
        crud = ArbitrationCaseCRUD(self.auth)
        data = self._case_payload(case)
        user = getattr(self.auth, "user", None)
        if user and data.get("user_id") is None:
            data["user_id"] = getattr(user, "id", None)
        if case.get("case_id"):
            existing = await crud.update_case(int(case["case_id"]), data)
            if existing:
                return existing
        return await crud.create_case(data)

    def _profile_defaults(self) -> dict[str, Any]:
        user = getattr(self.auth, "user", None)
        if user is None:
            return {}
        return {
            "applicant_name": getattr(user, "name", None),
            "applicant_phone": getattr(user, "mobile", None),
            "respondent_name": getattr(user, "company_name", None),
            "position_name": getattr(user, "position_name", None),
            "monthly_salary": getattr(user, "monthly_salary", None),
            "hire_date": getattr(user, "hire_date", None),
            "contract_type": getattr(user, "contract_type", None),
            "social_insurance": getattr(user, "social_insurance", None),
        }

    def _merge_case_data(self, data: ArbitrationDraftRequestSchema, profile: dict[str, Any]) -> dict[str, Any]:
        raw = data.model_dump()
        for key, value in profile.items():
            if raw.get(key) in (None, "") and value not in (None, ""):
                raw[key] = value
        raw["claims"] = raw.get("claims") or self._default_claims(raw)
        raw["evidence_items"] = raw.get("evidence_items") or self._default_evidence_items(raw)
        raw["case_title"] = raw.get("case_title") or self._default_case_title(raw)
        return raw

    @staticmethod
    def _default_case_title(case: dict[str, Any]) -> str:
        applicant = case.get("applicant_name") or "申请人"
        respondent = case.get("respondent_name") or "被申请人"
        return f"{applicant}与{respondent}劳动争议"

    async def _fetch_evidence_analyses(self, session_id: str | None) -> list[dict[str, Any]]:
        db = getattr(self.auth, "db", None)
        if not session_id or db is None:
            return []
        try:
            from app.plugin.module_ai.evidence.crud import EvidenceAnalysisCRUD

            return await EvidenceAnalysisCRUD(self.auth).get_analyzed_by_session(session_id=session_id)
        except Exception as exc:
            logger.warning(f"生成申请书时读取证据分析失败（非致命）: {exc}")
            return []

    @staticmethod
    def _default_claims(case: dict[str, Any]) -> list[str]:
        summary = str(case.get("dispute_summary") or "")
        claims = []
        if "工资" in summary or "欠薪" in summary or "拖欠" in summary:
            claims.append("请求裁决被申请人支付拖欠工资，具体金额以工资流水和实际欠付期间核算为准")
        if "解除" in summary or "辞退" in summary or "开除" in summary:
            claims.append("请求裁决被申请人依法支付解除劳动关系相关补偿或赔偿")
        if not claims:
            claims.append("请求裁决被申请人承担本案相关劳动争议责任，具体请求待补充")
        return claims

    @staticmethod
    def _default_evidence_items(case: dict[str, Any]) -> list[str]:
        evidence = ["劳动合同/入职材料", "工资流水/工资条", "考勤记录", "聊天记录/通知文件"]
        if case.get("social_insurance") is not None:
            evidence.append("社保缴纳记录")
        return evidence

    def _render_template(self, *, case: dict[str, Any], evidence_analyses: list[dict[str, Any]]) -> str:
        claims_text = self._numbered(case["claims"], suffix="；")
        evidence_text = self._render_evidence_items(case["evidence_items"], evidence_analyses)

        salary_text = self._format_salary(case.get("monthly_salary"))
        social_insurance = self._format_bool(case.get("social_insurance"))
        employment_fact = self._render_employment_fact(case, salary_text, social_insurance)
        today = date.today()

        return "\n".join(
            [
                "劳动人事争议仲裁申请书",
                "",
                f"申请人：{self._value(case.get('applicant_name'))}",
                f"性别：{self._value(case.get('applicant_gender'))}",
                f"身份证号：{self._value(case.get('applicant_id_no'))}",
                f"联系电话：{self._value(case.get('applicant_phone'))}",
                f"住址：{self._value(case.get('applicant_address'))}",
                "",
                f"被申请人：{self._value(case.get('respondent_name'))}",
                f"统一社会信用代码：{self._value(case.get('respondent_credit_code'))}",
                f"住所地：{self._value(case.get('respondent_address'))}",
                f"法定代表人：{self._value(case.get('respondent_legal_rep'))}",
                f"联系电话：{self._value(case.get('respondent_phone'))}",
                "",
                "仲裁请求：",
                "",
                claims_text,
                "",
                "事实与理由：",
                "",
                employment_fact,
                "",
                f"双方劳动关系存续期间发生以下争议：{case.get('dispute_summary')}",
                "",
                (
                    "申请人认为，被申请人的相关行为已经影响申请人的合法劳动权益。"
                    "现申请人依据已掌握的事实和证据材料，向贵委提出上述仲裁请求。"
                    "具体金额、期间和法律依据仍需结合证据材料进一步核对。"
                ),
                "",
                "证据目录：",
                "",
                evidence_text,
                "",
                "此致",
                "",
                self._value(case.get("arbitration_committee"), suffix="劳动人事争议仲裁委员会"),
                "",
                f"申请人：{self._value(case.get('applicant_name'))}",
                f"日期：{today.year}年{today.month}月{today.day}日",
            ]
        )

    def _render_employment_fact(self, case: dict[str, Any], salary_text: str, social_insurance: str) -> str:
        position = self._fact_value(case.get("position_name"))
        work_location = self._fact_value(case.get("work_location"))
        contract_type = self._fact_value(case.get("contract_type"))

        return (
            f"申请人于{self._value(case.get('hire_date'))}入职被申请人处工作，"
            f"{'岗位为' + position if position != MISSING else '岗位【待补充】'}，"
            f"{'工作地点为' + work_location if work_location != MISSING else '工作地点【待补充】'}，"
            f"双方约定工资标准为{salary_text}，"
            f"{'合同签订情况为' + contract_type if contract_type != MISSING else '合同签订情况【待补充】'}，"
            f"社会保险缴纳情况为{social_insurance}。"
        )

    @staticmethod
    def _numbered(items: list[str], *, suffix: str = "") -> str:
        return "\n".join(f"{index}. {item}{suffix}" for index, item in enumerate(items, start=1))

    def _render_evidence_items(self, items: list[str], analyses: list[dict[str, Any]]) -> str:
        lines = [f"{index}. {item}" for index, item in enumerate(items, start=1)]
        if lines:
            return "\n".join(lines)

        for offset, analysis in enumerate(analyses, start=1):
            file_name = analysis.get("file_name") or "已上传证据"
            summary = analysis.get("summary") or "已完成 AI 分析，提交前需人工复核"
            lines.append(f"{offset}. {file_name}：{summary}")
        return "\n".join(lines)

    @staticmethod
    def _value(value: Any, *, suffix: str = "") -> str:
        if value is None or value == "":
            return MISSING
        return f"{value}{suffix}" if suffix and str(value).endswith(suffix) is False else str(value)

    @staticmethod
    def _fact_value(value: Any) -> str:
        if value is None:
            return MISSING
        text = str(value).strip().strip("，。；;、 ")
        if not text:
            return MISSING
        if "工资标准及社保" in text or "社保缴纳约定" in text or "证明目的" in text:
            return MISSING
        return text

    @staticmethod
    def _format_salary(value: Any) -> str:
        if value is None or value == "":
            return MISSING
        try:
            amount = float(value)
        except (TypeError, ValueError):
            return str(value)
        if amount.is_integer():
            return f"{int(amount)}元/月"
        return f"{amount:.2f}元/月"

    @staticmethod
    def _format_bool(value: Any) -> str:
        if value is True:
            return "是"
        if value is False:
            return "否"
        return MISSING

    @staticmethod
    def _ai_configured() -> bool:
        api_key = settings.OPENAI_API_KEY.strip()
        return bool(api_key and api_key != "your_api_key" and settings.OPENAI_MODEL.strip())

    @staticmethod
    def _build_ai_prompt(template_content: str) -> str:
        return (
            "你是劳动仲裁文书辅助助手。请基于下方模板草稿润色成一份结构完整、表达正式的"
            "劳动人事争议仲裁申请书草稿。\n"
            "要求：不得编造身份证号、统一社会信用代码、日期、金额或证据；信息缺失保留【待补充】；"
            "不得承诺仲裁结果；只输出申请书正文。\n"
            "格式要求：输出纯文本正式文书，不要使用 Markdown，不要使用 **加粗**、# 标题、代码块、表格符号或项目符号。\n\n"
            f"{template_content}"
        )

    @staticmethod
    def _sanitize_ai_content(content: str) -> str:
        text = content.strip()
        text = re.sub(r"```(?:\w+)?\n?", "", text)
        text = text.replace("```", "")
        text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.MULTILINE)
        text = re.sub(r"\*\*([^*\n]+?)\*\*", r"\1", text)
        text = re.sub(r"__([^_\n]+?)__", r"\1", text)
        text = re.sub(r"^\s*[-*]\s+", "", text, flags=re.MULTILINE)
        text = re.sub(r"[ \t]+$", "", text, flags=re.MULTILINE)
        return text.strip()

    @staticmethod
    def _case_payload(case: dict[str, Any]) -> dict[str, Any]:
        keys = (
            "case_title",
            "session_id",
            "applicant_name",
            "applicant_gender",
            "applicant_id_no",
            "applicant_phone",
            "applicant_address",
            "respondent_name",
            "respondent_credit_code",
            "respondent_address",
            "respondent_legal_rep",
            "respondent_phone",
            "hire_date",
            "leave_date",
            "position_name",
            "work_location",
            "monthly_salary",
            "contract_type",
            "social_insurance",
            "dispute_summary",
            "claims",
            "evidence_items",
        )
        payload = {key: case.get(key) for key in keys}
        payload["user_id"] = case.get("user_id")
        return payload

    def _case_to_dict(self, case: Any) -> dict[str, Any]:
        return {
            "id": case.id,
            "case_title": case.case_title,
            "respondent_name": case.respondent_name,
            "applicant_name": case.applicant_name,
            "dispute_summary": case.dispute_summary,
            "session_id": case.session_id,
            "created_time": str(case.created_time) if case.created_time else None,
            "updated_time": str(case.updated_time) if case.updated_time else None,
        }

    def _draft_to_dict(self, draft: Any) -> dict[str, Any]:
        return {
            "id": draft.id,
            "case_id": draft.case_id,
            "session_id": draft.session_id,
            "title": draft.title,
            "review_status": draft.review_status,
            "used_ai": draft.used_ai,
            "created_time": str(draft.created_time) if draft.created_time else None,
            "updated_time": str(draft.updated_time) if draft.updated_time else None,
        }

    @staticmethod
    def _json_safe(value: Any) -> Any:
        if isinstance(value, Decimal):
            return float(value)
        if isinstance(value, dict):
            return {key: ArbitrationDraftService._json_safe(item) for key, item in value.items()}
        if isinstance(value, list):
            return [ArbitrationDraftService._json_safe(item) for item in value]
        return value

    @staticmethod
    def _build_docx(title: str, content: str) -> bytes:
        from docx import Document

        document = Document()
        document.add_heading(title, level=0)
        for block in content.split("\n"):
            document.add_paragraph(block)
        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _build_pdf(title: str, content: str) -> bytes:
        lines = [title, "", *content.splitlines()]
        objects: list[bytes] = []

        def pdf_text(text: str) -> str:
            return "<FEFF" + text.encode("utf-16-be").hex().upper() + ">"

        content_lines = ["BT", "/F1 11 Tf", "50 790 Td", "14 TL"]
        for line in lines[:48]:
            content_lines.append(f"{pdf_text(line[:80])} Tj")
            content_lines.append("T*")
        content_lines.append("ET")
        stream = "\n".join(content_lines).encode()

        objects.append(b"<< /Type /Catalog /Pages 2 0 R >>")
        objects.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        )
        objects.append(
            b"<< /Type /Font /Subtype /Type0 /BaseFont /STSong-Light "
            b"/Encoding /UniGB-UCS2-H /DescendantFonts [6 0 R] >>"
        )
        objects.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        objects.append(
            b"<< /Type /Font /Subtype /CIDFontType0 /BaseFont /STSong-Light "
            b"/CIDSystemInfo << /Registry (Adobe) /Ordering (GB1) /Supplement 2 >> >>"
        )

        output = io.BytesIO()
        output.write(b"%PDF-1.4\n")
        offsets = [0]
        for index, obj in enumerate(objects, start=1):
            offsets.append(output.tell())
            output.write(f"{index} 0 obj\n".encode())
            output.write(obj)
            output.write(b"\nendobj\n")
        xref = output.tell()
        output.write(f"xref\n0 {len(objects) + 1}\n".encode())
        output.write(b"0000000000 65535 f \n")
        for offset in offsets[1:]:
            output.write(f"{offset:010d} 00000 n \n".encode())
        output.write(
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
        )
        return output.getvalue()
