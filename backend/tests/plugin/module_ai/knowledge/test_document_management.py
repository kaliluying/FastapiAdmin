import io
from unittest.mock import AsyncMock

import pytest
import test_document_processing
from docx import Document
from fastapi import BackgroundTasks, UploadFile
from sqlalchemy import update

from app.core.exceptions import CustomException
from app.plugin.module_ai.knowledge import document_lifecycle as service_module
from app.plugin.module_ai.knowledge import extractors
from app.plugin.module_ai.knowledge.model import KnowledgeDocumentModel

processing = test_document_processing.processing
document_snapshot = test_document_processing.document_snapshot


async def test_replacement_activates_file_and_chunks_together(processing):
    background = BackgroundTasks()
    await processing.lifecycle.replace_document(processing.document_id, UploadFile(filename="replacement.txt", file=io.BytesIO("纠正后的正文".encode())), background)
    staged, chunks = await document_snapshot(processing)
    assert staged.pending_file and staged.file_name == "guide.txt"
    assert chunks[0].content == "原有内容"
    path, filename = await processing.service.document_file(staged.id)
    assert filename == "guide.txt" and path.read_text() == "可检索的文档内容"
    await background()
    document, chunks = await document_snapshot(processing)
    assert document.index_status == "success" and document.pending_file is None
    assert document.file_name == "replacement.txt"
    assert chunks[0].content == "纠正后的正文"
    assert not processing.path.exists()
    path, _name = await processing.service.document_file(staged.id)
    assert path.read_text() == "纠正后的正文"


async def test_failed_replacement_preserves_old_file_chunks_and_retry(processing):
    processing.embeddings.embed_texts.side_effect = RuntimeError("provider secret")
    background = BackgroundTasks()
    await processing.lifecycle.replace_document(processing.document_id, UploadFile(filename="fixed.txt", file=io.BytesIO(b"new text")), background)
    await background()
    document, chunks = await document_snapshot(processing)
    assert document.index_status == "failed" and document.pending_file
    assert document.file_path == processing.path.name and processing.path.exists()
    assert chunks[0].chroma_id == "old-chunk"
    assert processing.vectors == processing.keywords == {"old-chunk": "原有内容"}
    processing.embeddings.embed_texts.side_effect = None
    await processing.lifecycle.index_document(document.id)
    document, chunks = await document_snapshot(processing)
    assert document.index_status == "success" and chunks[0].content == "new text"


async def test_recovery_uses_persisted_pending_replacement(processing):
    await processing.lifecycle.replace_document(processing.document_id, UploadFile(filename="fixed.txt", file=io.BytesIO(b"restart text")), BackgroundTasks())
    await service_module.KnowledgeDocumentLifecycle.recover_pending_documents()
    document, chunks = await document_snapshot(processing)
    assert document.index_status == "success" and document.pending_file is None
    assert chunks[0].content == "restart text"


async def test_conflicting_replace_does_not_leak_upload(processing):
    await processing.db.execute(update(KnowledgeDocumentModel).where(KnowledgeDocumentModel.id == processing.document_id).values(index_status="indexing"))
    await processing.db.commit()
    before = set(processing.path.parent.iterdir())
    with pytest.raises(CustomException) as error:
        await processing.lifecycle.replace_document(processing.document_id, UploadFile(filename="fixed.txt", file=io.BytesIO(b"text")), BackgroundTasks())
    assert error.value.status_code == 409
    assert set(processing.path.parent.iterdir()) == before


async def test_download_rejects_path_traversal_and_symlinks(processing, tmp_path):
    outside = tmp_path.parent / "secret-document.txt"
    outside.write_text("private")
    link = tmp_path / "link.txt"
    link.symlink_to(outside)
    for path in ("../secret-document.txt", str(outside), link.name):
        await processing.db.execute(update(KnowledgeDocumentModel).where(KnowledgeDocumentModel.id == processing.document_id).values(file_path=path))
        await processing.db.commit()
        with pytest.raises(CustomException, match="路径非法"):
            await processing.service.document_file(processing.document_id)


async def test_download_and_chunks_require_access_before_reading(processing, monkeypatch):
    monkeypatch.setattr(service_module.KnowledgeDocumentCRUD, "get_or_404", AsyncMock(side_effect=CustomException(msg="denied", status_code=403)))
    for call in (
        processing.service.document_file(processing.document_id),
        processing.service.page_chunks(processing.document_id, page_no=1, page_size=10),
        processing.service.document_content(processing.document_id),
    ):
        with pytest.raises(CustomException, match="denied"):
            await call


async def test_chunk_page_and_content_correction(processing):
    page = await processing.service.page_chunks(processing.document_id, page_no=1, page_size=1)
    assert page["total"] == 1 and page["items"][0]["content"] == "原有内容"
    assert await processing.service.document_content(processing.document_id) == "可检索的文档内容"
    with pytest.raises(CustomException, match="不能为空"):
        await processing.lifecycle.update_document_content(processing.document_id, "  ", BackgroundTasks())
    background = BackgroundTasks()
    await processing.lifecycle.update_document_content(processing.document_id, "人工纠错正文", background)
    await background()
    document, chunks = await document_snapshot(processing)
    assert document.file_type == "txt" and chunks[0].content == "人工纠错正文"


def test_docx_preserves_paragraph_table_order_and_merged_cells(tmp_path):
    document = Document()
    document.add_paragraph("前言")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "品名"
    table.cell(0, 1).text = "报价"
    table.cell(1, 0).merge(table.cell(1, 1)).text = "合并说明"
    document.add_paragraph("结语")
    path = tmp_path / "quote.docx"
    document.save(path)
    assert extractors._extract_text_sync(path) == "前言\n品名\t报价\n合并说明\n结语"


def test_scanned_pdf_reports_missing_ocr_dependency(monkeypatch, tmp_path):
    from PIL import Image, ImageDraw

    image = Image.new("RGB", (600, 160), "white")
    ImageDraw.Draw(image).text((20, 40), "Scanned invoice", fill="black")
    path = tmp_path / "scan.pdf"
    image.save(path, "PDF")
    monkeypatch.setattr(extractors.shutil, "which", lambda _name: None)
    with pytest.raises(extractors.DocumentExtractionError, match="Poppler"):
        extractors._extract_text_sync(path)


def test_ocr_cli_uses_bounded_page_and_available_languages(monkeypatch, tmp_path):
    from types import SimpleNamespace

    calls = []
    monkeypatch.setattr(extractors.shutil, "which", lambda name: name)

    def run(args, **kwargs):
        calls.append((args, kwargs))
        return SimpleNamespace(stdout="eng\nchi_sim\n" if "--list-langs" in args else "已识别内容")

    monkeypatch.setattr(extractors.subprocess, "run", run)
    assert extractors._ocr_pdf_page(tmp_path / "scan.pdf", 3) == "已识别内容"
    assert calls[1][0][1:9] == ["-f", "3", "-l", "3", "-singlefile", "-scale-to", "2400", "-png"]
    assert calls[2][0][-1] == "chi_sim+eng"
    assert all(call[1]["timeout"] <= 60 for call in calls)


async def test_replacing_a_failed_pending_file_cleans_discarded_upload(processing):
    processing.embeddings.embed_texts.side_effect = RuntimeError("temporary")
    first = BackgroundTasks()
    await processing.lifecycle.replace_document(processing.document_id, UploadFile(filename="first.txt", file=io.BytesIO(b"first")), first)
    await first()
    failed, _chunks = await document_snapshot(processing)
    discarded = processing.path.parent / failed.pending_file["file_path"]
    second = BackgroundTasks()
    await processing.lifecycle.replace_document(processing.document_id, UploadFile(filename="second.txt", file=io.BytesIO(b"second")), second)
    processing.embeddings.embed_texts.side_effect = None
    await second()
    document, chunks = await document_snapshot(processing)
    assert document.index_status == "success" and chunks[0].content == "second"
    assert not discarded.exists() and not processing.path.exists()


async def test_final_commit_failure_retains_original_file_and_pending_replacement(processing, monkeypatch):
    await processing.lifecycle.replace_document(processing.document_id, UploadFile(filename="fixed.txt", file=io.BytesIO(b"fixed")), BackgroundTasks())
    commit = processing.db.commit
    attempts = 0

    async def fail_commit():
        nonlocal attempts
        attempts += 1
        if attempts == 5:
            raise RuntimeError("commit failure")
        await commit()

    monkeypatch.setattr(processing.db, "commit", fail_commit)
    with pytest.raises(CustomException):
        await processing.lifecycle.index_document(processing.document_id)
    document, chunks = await document_snapshot(processing)
    assert document.file_path == processing.path.name and document.pending_file
    assert processing.path.exists() and chunks[0].chroma_id == "old-chunk"
    assert processing.vectors == processing.keywords == {"old-chunk": "原有内容"}


async def test_deleting_failed_replacement_cleans_original_and_pending_files(processing):
    processing.embeddings.embed_texts.side_effect = RuntimeError("temporary")
    background = BackgroundTasks()
    await processing.lifecycle.replace_document(processing.document_id, UploadFile(filename="pending.txt", file=io.BytesIO(b"pending")), background)
    await background()
    failed, _chunks = await document_snapshot(processing)
    pending_path = processing.path.parent / failed.pending_file["file_path"]
    processing.store.delete_document.side_effect = None
    processing.bm25.delete_by_document.side_effect = None
    await processing.lifecycle.delete_document([processing.document_id])
    assert not processing.path.exists() and not pending_path.exists()
