import asyncio
import json
from contextlib import asynccontextmanager, suppress
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock
from uuid import uuid4

import pytest
from docx import Document
from fastapi import UploadFile
from pydantic import ValidationError
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.v1.module_platform.menu.model import MenuModel
from app.api.v1.module_system.role.model import RoleModel
from app.api.v1.module_system.user.model import UserModel
from app.config.setting import settings
from app.core.database import async_db_session
from app.core.exceptions import CustomException
from app.plugin.module_ai.chat import service, ws
from app.plugin.module_ai.chat.crud import ChatSession, ChatSessionCRUD
from app.plugin.module_ai.chat.hybrid_retriever import KnowledgeBaseChatRetriever
from app.plugin.module_ai.chat.model import ChatSessionModel
from app.plugin.module_ai.chat.rag import MAX_CONTEXT_CHARS, MAX_DOCUMENT_CHARS, RagChatChain, RagDocument, RagPromptBuilder
from app.plugin.module_ai.chat.schema import ChatQuerySchema, ChatSessionCreateSchema
from app.plugin.module_ai.knowledge import public
from app.plugin.module_ai.knowledge.model import KnowledgeBaseModel, KnowledgeChunkModel, KnowledgeDocumentModel
from app.plugin.module_ai.knowledge.public import KnowledgeSearchResult
from app.utils.hash_bcrpy_util import PwdUtil


@pytest.fixture
async def chat_auth():
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    try:
        async with engine.begin() as connection:
            await connection.run_sync(ChatSessionModel.__table__.create)
        async with async_sessionmaker(engine, expire_on_commit=False)() as database:
            yield SimpleNamespace(db=database, user=SimpleNamespace(id=1, username="admin", is_superuser=True))
    finally:
        await engine.dispose()


@pytest.fixture
def configured_chat(monkeypatch):
    monkeypatch.setattr(service, "load_runtime_chat_model_config", AsyncMock())
    monkeypatch.setattr(service.ChatService, "_validate_ai_config", staticmethod(lambda: None))
    monkeypatch.setattr(service.ChatService, "_trigger_memory_extraction", lambda *args, **kwargs: None)


@pytest.fixture(scope="module")
def websocket_scope_data(ai_client):
    assert settings.DATABASE_TYPE == "sqlite"

    async def seed():
        suffix = uuid4().hex[:8]
        async with async_db_session() as database:
            menus = [MenuModel(name=f"WS Scope {suffix} {index}", type=3, permission=permission, status=0) for index, permission in enumerate(["module_ai:chat:ws", "module_ai:session:detail"])]
            role = RoleModel(name=f"WS Own {suffix}", code=f"WS_OWN_{suffix}", status=0, data_scope=1, menus=menus)
            password = PwdUtil.hash_password("isolated-scope-password")
            owner = UserModel(username=f"ws_owner_{suffix}", name="WS Scope Owner", password=password, is_superuser=False, status=0, roles=[role])
            other = UserModel(username=f"ws_other_{suffix}", name="WS Scope Other", password=password, is_superuser=False, status=0, roles=[role])
            database.add_all([owner, other])
            await database.flush()
            bases = []
            for user, marker in [(owner, f"OWN_KB_BODY_{suffix}"), (other, f"FOREIGN_KB_SECRET_{suffix}")]:
                base = KnowledgeBaseModel(name=f"WS Base {user.username}", created_id=user.id, is_enabled=True)
                database.add(base)
                await database.flush()
                document = KnowledgeDocumentModel(knowledge_base_id=base.id, file_name=f"{marker}.md", file_type=".md", parse_status="parsed", index_status="indexed", created_id=user.id)
                database.add(document)
                await database.flush()
                database.add(KnowledgeChunkModel(knowledge_base_id=base.id, document_id=document.id, chunk_index=0, content=marker, chroma_id=f"ws-scope-{suffix}-{user.id}", created_id=user.id))
                bases.append({"id": base.id, "document_id": document.id, "marker": marker})
            await database.commit()
            return {"username": owner.username, "user_id": owner.id, "own": bases[0], "foreign": bases[1]}

    return ai_client.portal.call(seed)


class QueueWebSocket:
    query_params = {"ticket": "test-ticket"}
    client = "isolated-test"

    def __init__(self):
        self.incoming = asyncio.Queue()
        self.outgoing = asyncio.Queue()
        self.accepted = False

    async def accept(self):
        self.accepted = True

    async def receive_text(self):
        message = await self.incoming.get()
        if message is None:
            raise RuntimeError("test disconnect")
        return json.dumps(message, ensure_ascii=False)

    async def send_text(self, message):
        await self.outgoing.put(message)

    async def send(self, message):
        await self.incoming.put(message)

    async def event(self):
        return json.loads(await asyncio.wait_for(self.outgoing.get(), timeout=2))


@asynccontextmanager
async def running_socket(monkeypatch, chat_query):
    opened = []
    exited = []
    auth = SimpleNamespace(user=SimpleNamespace(id=1, username="admin", is_superuser=True))

    @asynccontextmanager
    async def database_session():
        database = SimpleNamespace(rollback=AsyncMock())
        opened.append(database)
        try:
            yield database
        finally:
            await database.rollback()
            exited.append(database)

    monkeypatch.setattr(ws, "async_db_session", database_session)
    monkeypatch.setattr(ws, "_resolve_ws_auth", AsyncMock(return_value=auth))
    monkeypatch.setattr(ws.ChatService, "chat_query", chat_query)
    websocket = QueueWebSocket()
    task = asyncio.create_task(ws.websocket_chat_controller(websocket))
    try:
        yield websocket, opened, exited
    finally:
        await websocket.incoming.put(None)
        try:
            await asyncio.wait_for(task, timeout=2)
        finally:
            if not task.done():
                task.cancel()
                with suppress(asyncio.CancelledError):
                    await task


async def test_structured_query_streams_real_rag_and_persists_citations(chat_auth, configured_chat, monkeypatch):
    captured = {}

    class Retriever:
        async def retrieve(self, **kwargs):
            captured["query"] = kwargs
            return [RagDocument(content="attachment policy approved", metadata={"name": "policy.md", "private_path": "/private/server-file"})]

    class Model:
        async def stream(self, prompt):
            captured["prompt"] = prompt
            yield "policy "
            yield "approved"

    def create_chain(db, auth):
        assert db is chat_auth.db
        return RagChatChain(retriever=Retriever(), prompt_builder=RagPromptBuilder(), chat_model=Model(), db=db)

    monkeypatch.setattr(service, "create_rag_chain", create_chain)
    session = await ChatSessionCRUD(chat_auth).create_crud(ChatSessionCreateSchema(title="isolated citations"))
    await chat_auth.db.commit()
    files = [{"name": "policy.md", "content": "attachment policy approved"}]
    events = [
        event async for event in service.ChatService(chat_auth).chat_query(ChatQuerySchema(message="policy", session_id=session.session_id, request_id="request-1", files=files), structured=True)
    ]

    assert [event["type"] for event in events] == ["stage", "citations", "stage", "chunk", "chunk", "done"]
    assert [event["stage"] for event in events if event["type"] == "stage"] == ["retrieving", "generating"]
    assert events[-1] == {"type": "done", "session_id": session.session_id}
    assert captured["query"]["files"] == files
    assert "attachment policy approved" in captured["prompt"]
    assert events[1]["citations"] == [{"id": "1", "title": "policy.md", "snippet": "attachment policy approved"}]
    stored = await ChatSessionCRUD(chat_auth).get_by_id_crud(session.session_id)
    assert stored.runs[0]["messages"][1]["content"] == "policy approved"
    assert stored.runs[0]["messages"][1]["citations"] == events[1]["citations"]
    assert "private_path" not in json.dumps(stored.runs)


@pytest.mark.parametrize("failure", ["empty", "provider", "persistence", "commit"])
async def test_structured_query_failures_never_emit_done_or_expose_internal_error(configured_chat, monkeypatch, failure):
    database = SimpleNamespace(commit=AsyncMock(), rollback=AsyncMock())
    if failure == "commit":
        database.commit.side_effect = RuntimeError("private-database-connection")

    class Crud:
        def __init__(self, auth):
            self.db = auth.db

        async def append_run_crud(self, **kwargs):
            return failure != "persistence"

    class Chain:
        async def astream(self, **kwargs):
            assert kwargs["include_events"] is True
            if failure == "empty":
                return
            yield {"type": "citations", "citations": []}
            yield {"type": "stage", "stage": "generating"}
            if failure == "provider":
                raise RuntimeError("private-provider-api-key")
            yield "partial answer"

    monkeypatch.setattr(service, "ChatSessionCRUD", Crud)
    monkeypatch.setattr(service, "create_rag_chain", lambda **kwargs: Chain())
    auth = SimpleNamespace(db=database, user=SimpleNamespace(id=1))
    events = [event async for event in service.ChatService(auth).chat_query(ChatQuerySchema(message="hello", session_id="isolated", request_id="failed-request"), structured=True)]

    assert events[-1]["type"] == "error"
    assert all(event["type"] != "done" for event in events)
    assert "private-" not in json.dumps(events)
    if failure != "empty":
        database.rollback.assert_awaited_once()


async def test_structured_configuration_error_is_terminal(configured_chat, monkeypatch):
    monkeypatch.setattr(service.ChatService, "_validate_ai_config", staticmethod(lambda: "AI 服务未配置有效 API Key"))
    events = [event async for event in service.ChatService(SimpleNamespace()).chat_query(ChatQuerySchema(message="hello", request_id="configuration"), structured=True)]
    assert events == [{"type": "error", "message": "AI 服务未配置有效 API Key"}]


@pytest.mark.parametrize("knowledge_base_ids", [[], [7]])
async def test_attachments_are_context_without_keyword_match_and_with_selected_knowledge(knowledge_base_ids):
    search = SimpleNamespace(
        search=AsyncMock(return_value=[KnowledgeSearchResult(chunk_id="stored-1", content="知识库参考正文", metadata={"knowledge_base_id": 7, "document_id": 9, "file_name": "stored.md"}, score=1.0)])
    )
    retriever = KnowledgeBaseChatRetriever(mode="bm25", knowledge_search=search)
    files = [{"name": "temporary.md", "content": "只含附件独特标记XYZ"}]
    documents = await retriever.retrieve(query="请总结这些材料", user_id="1", scope_id="1", session_id=None, knowledge_base_ids=knowledge_base_ids, files=files)

    assert documents[0].content == files[0]["content"]
    assert documents[0].metadata == {"source": "uploaded-file", "name": "temporary.md"}
    assert len(documents) == (2 if knowledge_base_ids else 1)
    if knowledge_base_ids:
        search.search.assert_awaited_once_with(query="请总结这些材料", knowledge_base_ids=[7])
        assert documents[1].metadata["knowledge_base_id"] == 7
        assert documents[1].metadata["bm25_score"] == 1.0
    else:
        search.search.assert_not_awaited()


async def test_legacy_crud_preserves_transcript_shape_without_citations(chat_auth):
    crud = ChatSessionCRUD(chat_auth)
    session = await crud.create_crud(ChatSessionCreateSchema(title="legacy"))
    assert await crud.append_run_crud(session_id=session.session_id, message="hello", response="legacy answer")
    await chat_auth.db.commit()
    stored = await crud.get_by_id_crud(session.session_id)
    assert stored.runs[0]["messages"] == [{"role": "user", "content": "hello"}, {"role": "assistant", "content": "legacy answer"}]


async def test_five_long_attachments_only_cite_the_context_actually_given_to_model():
    captured = {}
    files = [{"name": f"attachment-{index + 1}.txt", "content": f"独特开头-{index + 1}-" + character * 15_980} for index, character in enumerate("壹贰叁肆伍")]

    class Model:
        async def stream(self, prompt):
            captured["prompt"] = prompt
            yield "bounded answer"

    retriever = KnowledgeBaseChatRetriever(mode="bm25", knowledge_search=SimpleNamespace(search=AsyncMock(return_value=[])))
    chain = RagChatChain(retriever=retriever, prompt_builder=RagPromptBuilder(), chat_model=Model())
    events = [event async for event in chain.astream(message="请总结材料", user_id="1", scope_id="1", session_id=None, files=files, include_events=True)]
    citations = events[0]["citations"]
    context = captured["prompt"].split("检索上下文:\n", 1)[1].split("\n\n用户问题:\n", 1)[0]

    assert len(context) <= MAX_CONTEXT_CHARS
    assert [citation["title"] for citation in citations] == [file["name"] for file in files[:3]]
    assert [citation["id"] for citation in citations] == ["1", "2", "3"]
    assert all(citation["snippet"] in context for citation in citations)
    assert len(citations[-1]["snippet"]) < MAX_DOCUMENT_CHARS
    assert files[2]["content"][:MAX_DOCUMENT_CHARS] not in context
    for file in files[3:]:
        assert file["name"] not in context
        assert file["content"][:20] not in context


async def test_websocket_adds_request_id_to_all_structured_events(monkeypatch):
    events = [
        {"type": "stage", "stage": "retrieving"},
        {"type": "citations", "citations": [{"id": "1", "title": "文档", "snippet": "正文"}]},
        {"type": "stage", "stage": "generating"},
        {"type": "chunk", "content": "回答"},
        {"type": "done", "session_id": "session-1"},
    ]

    async def chat_query(self, query, structured=False):
        assert structured is True
        assert query.files[0]["content"] == "正文"
        for event in events:
            yield event

    async with running_socket(monkeypatch, chat_query) as (websocket, opened, exited):
        await websocket.send({"message": "hello", "request_id": "request-1", "files": [{"name": "doc.txt", "content": "正文"}]})
        received = [await websocket.event() for _ in events]
        assert received == [{**event, "request_id": "request-1"} for event in events]
        assert websocket.accepted
    assert len(opened) == len(exited) == 2


@pytest.mark.parametrize("disconnect", [False, True])
async def test_websocket_cancellation_releases_generator_and_database(monkeypatch, disconnect):
    released = asyncio.Event()
    started = asyncio.Event()

    async def chat_query(self, query, structured=False):
        try:
            if query.message == "after cancellation":
                yield {"type": "done", "session_id": "after-cancel"}
                return
            started.set()
            yield {"type": "chunk", "content": "partial"}
            await asyncio.Event().wait()
        finally:
            released.set()

    async with running_socket(monkeypatch, chat_query) as (websocket, opened, exited):
        await websocket.send({"message": "hello", "request_id": "cancel-1"})
        assert (await websocket.event())["type"] == "chunk"
        await asyncio.wait_for(started.wait(), timeout=2)
        await websocket.send(None if disconnect else {"type": "cancel", "request_id": "cancel-1"})
        assert await websocket.event() == {"type": "cancelled", "request_id": "cancel-1"}
        await asyncio.wait_for(released.wait(), timeout=2)
        if not disconnect:
            await websocket.send({"message": "after cancellation", "request_id": "cancel-2"})
            assert await websocket.event() == {"type": "done", "session_id": "after-cancel", "request_id": "cancel-2"}
    assert len(opened) == len(exited) == (2 if disconnect else 3)
    for database in opened:
        database.rollback.assert_awaited_once()


async def test_websocket_wrong_cancel_does_not_cancel_and_busy_request_is_rejected(monkeypatch):
    released = asyncio.Event()

    async def chat_query(self, query, structured=False):
        try:
            yield {"type": "stage", "stage": "generating"}
            await asyncio.Event().wait()
        finally:
            released.set()

    async with running_socket(monkeypatch, chat_query) as (websocket, _, _):
        await websocket.send({"message": "first", "request_id": "first"})
        assert (await websocket.event())["request_id"] == "first"
        await websocket.send({"type": "cancel", "request_id": "other"})
        await websocket.send({"message": "second", "request_id": "second"})
        event = await websocket.event()
        assert event["type"] == "error" and event["request_id"] == "second"
        assert not released.is_set()
        await websocket.send({"type": "cancel", "request_id": "first"})
        assert (await websocket.event())["type"] == "cancelled"


async def test_websocket_recovers_after_validation_and_runtime_errors(monkeypatch):
    async def chat_query(self, query, structured=False):
        if query.message == "fail":
            raise RuntimeError("private-provider-secret")
        yield {"type": "done", "session_id": "session-1"}

    async with running_socket(monkeypatch, chat_query) as (websocket, _, _):
        await websocket.send({"message": "", "request_id": "invalid"})
        invalid = await websocket.event()
        assert invalid["type"] == "error" and invalid["request_id"] == "invalid"
        await websocket.send({"message": "fail", "request_id": "failed"})
        failure = await websocket.event()
        assert failure["type"] == "error" and failure["request_id"] == "failed"
        assert "private-" not in json.dumps(failure)
        await websocket.send({"message": "retry", "request_id": "retried"})
        assert await websocket.event() == {"type": "done", "session_id": "session-1", "request_id": "retried"}


async def test_websocket_legacy_client_receives_plain_text(monkeypatch):
    async def chat_query(self, query):
        assert query.request_id is None
        yield "legacy part "
        yield "two"

    async with running_socket(monkeypatch, chat_query) as (websocket, _, _):
        await websocket.send({"message": "legacy"})
        assert await asyncio.wait_for(websocket.outgoing.get(), timeout=2) == "legacy part "
        assert await asyncio.wait_for(websocket.outgoing.get(), timeout=2) == "two"


def test_authenticated_attachment_websocket_and_history_flow(ai_client, configured_chat, monkeypatch):
    class Model:
        async def stream(self, prompt):
            assert "材料独特标记XYZ" in prompt
            yield "测试回答"

    def create_chain(db, auth):
        retriever = KnowledgeBaseChatRetriever(mode="bm25", knowledge_search=SimpleNamespace(search=AsyncMock(return_value=[])))
        return RagChatChain(retriever=retriever, prompt_builder=RagPromptBuilder(), chat_model=Model(), db=db)

    monkeypatch.setattr(service, "create_rag_chain", create_chain)
    unauthorized = ai_client.post("/ai/chat/attachment", files={"file": ("material.txt", "材料独特标记XYZ".encode(), "text/plain")})
    assert unauthorized.status_code in {401, 403}
    login = ai_client.post("/system/auth/login", data={"username": "admin", "password": "admin123"})
    assert login.status_code == 200
    headers = {"Authorization": f"Bearer {login.json()['data']['access_token']}"}
    upload = ai_client.post("/ai/chat/attachment", headers=headers, files={"file": ("material.txt", "材料独特标记XYZ".encode(), "text/plain")})
    assert upload.status_code == 200
    attachment = upload.json()["data"]
    created = ai_client.post("/ai/chat/create", headers=headers, json={"title": "isolated upload flow"})
    assert created.status_code == 200
    session_id = created.json()["data"]["session_id"]
    ticket = ai_client.post("/system/auth/ws-ticket", headers=headers)
    assert ticket.status_code == 200
    with ai_client.websocket_connect(f"/ai/chat/ws?ticket={ticket.json()['data']['ticket']}") as websocket:
        websocket.send_json({"message": "请总结材料", "session_id": session_id, "request_id": "api-flow", "files": [attachment]})
        events = []
        for _ in range(8):
            event = websocket.receive_json()
            events.append(event)
            if event["type"] in {"done", "error"}:
                break
    assert [event["type"] for event in events] == ["stage", "citations", "stage", "chunk", "done"]
    assert all(event["request_id"] == "api-flow" for event in events)
    detail = ai_client.get(f"/ai/chat/detail/{session_id}", headers=headers)
    assert detail.status_code == 200
    assistant = detail.json()["data"]["messages"][-1]
    assert assistant["content"] == "测试回答"
    assert assistant["citations"] == [{"id": "1", "title": "material.txt", "snippet": "材料独特标记XYZ"}]


@pytest.mark.parametrize("scenario", ["cross_user_denied", "own_allowed", "superuser_allowed"])
def test_real_websocket_ticket_auth_enforces_knowledge_creator_scope(ai_client, websocket_scope_data, configured_chat, monkeypatch, scenario):
    searched_ids = []
    prompts = []
    derived_auth = []
    base = websocket_scope_data["own" if scenario == "own_allowed" else "foreign"]

    class Search:
        def __init__(self, database):
            self.database = database

        async def search(self, *, query, knowledge_base_ids, top_k=None):
            searched_ids.append(knowledge_base_ids)
            result = await self.database.execute(
                select(KnowledgeChunkModel, KnowledgeDocumentModel.file_name)
                .join(KnowledgeDocumentModel, KnowledgeDocumentModel.id == KnowledgeChunkModel.document_id)
                .where(KnowledgeChunkModel.knowledge_base_id.in_(knowledge_base_ids))
            )
            return [
                KnowledgeSearchResult(
                    chunk_id=chunk.chroma_id, content=chunk.content, metadata={"knowledge_base_id": chunk.knowledge_base_id, "document_id": chunk.document_id, "file_name": filename}, score=1.0
                )
                for chunk, filename in result.all()
            ]

    class Model:
        async def stream(self, prompt):
            prompts.append(prompt)
            yield "本地隔离回答"

    def create_chain(db, auth):
        derived_auth.append((auth.check_data_scope, auth.user.is_superuser))
        retriever = KnowledgeBaseChatRetriever(mode="bm25", knowledge_search=Search(db))
        return RagChatChain(retriever=retriever, prompt_builder=RagPromptBuilder(), chat_model=Model(), db=db)

    monkeypatch.setattr(service, "create_rag_chain", create_chain)
    username = "admin" if scenario == "superuser_allowed" else websocket_scope_data["username"]
    password = "admin123" if scenario == "superuser_allowed" else "isolated-scope-password"
    login = ai_client.post("/system/auth/login", data={"username": username, "password": password})
    assert login.status_code == 200
    headers = {"Authorization": f"Bearer {login.json()['data']['access_token']}"}
    ticket = ai_client.post("/system/auth/ws-ticket", headers=headers)
    assert ticket.status_code == 200
    with ai_client.websocket_connect(f"/ai/chat/ws?ticket={ticket.json()['data']['ticket']}") as websocket:
        websocket.send_json({"message": "请总结知识库资料", "knowledge_base_ids": [base["id"]], "request_id": scenario})
        events = []
        for _ in range(8):
            event = websocket.receive_json()
            events.append(event)
            if event["type"] in {"done", "error"}:
                break
    assert all(event["request_id"] == scenario for event in events)
    if scenario == "cross_user_denied":
        assert [event["type"] for event in events] == ["error"]
        assert base["marker"] not in json.dumps(events)
        assert searched_ids == []
        assert prompts == []
        assert derived_auth == []
    else:
        assert [event["type"] for event in events] == ["stage", "citations", "stage", "chunk", "done"]
        assert searched_ids == [[base["id"]]]
        assert derived_auth == [(True, scenario == "superuser_allowed")]
        assert len(prompts) == 1 and base["marker"] in prompts[0]
        assert events[1]["citations"] == [{"id": "1", "title": f"{base['marker']}.md", "snippet": base["marker"], "knowledge_base_id": base["id"], "document_id": base["document_id"]}]
        detail = ai_client.get(f"/ai/chat/detail/{events[-1]['session_id']}", headers=headers)
        assert detail.status_code == 200
        assert detail.json()["data"]["messages"][-1]["citations"] == events[1]["citations"]


def pdf_bytes(with_text=True):
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    if with_text:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 72 720 Td (attachment policy approved) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.mark.parametrize("suffix", [".txt", ".md", ".docx", ".pdf"])
async def test_attachment_extracts_real_supported_formats(suffix):
    content = b"attachment policy approved"
    if suffix == ".docx":
        document = Document()
        document.add_paragraph(content.decode())
        output = BytesIO()
        document.save(output)
        content = output.getvalue()
    elif suffix == ".pdf":
        content = pdf_bytes()
    upload = UploadFile(filename=f"../../policy{suffix}", file=BytesIO(content))
    result = await public.extract_chat_attachment(upload)
    assert result == {"name": f"policy{suffix}", "size": len(content), "type": suffix, "content": "attachment policy approved", "truncated": False}
    assert upload.file.closed


@pytest.mark.parametrize(
    "filename,content,expected",
    [
        ("policy.exe", b"not supported", "仅支持"),
        ("policy.txt", b" ", "没有可读取"),
        ("policy.pdf", b"broken private filename", "解析失败"),
        ("policy.docx", b"broken private filename", "解析失败"),
        ("policy.pdf", pdf_bytes(False), "没有可读取"),
        ("policy.txt", b"x" * (10 * 1024 * 1024 + 1), "10MB"),
    ],
)
async def test_attachment_rejects_invalid_empty_and_oversized_content(filename, content, expected):
    upload = UploadFile(filename=filename, file=BytesIO(content))
    try:
        with pytest.raises(CustomException, match=expected) as failure:
            await public.extract_chat_attachment(upload)
        assert "private filename" not in str(failure.value)
    finally:
        await upload.close()


async def test_attachment_truncates_context_and_cleans_temporary_file(monkeypatch):
    original = public.extract_text
    paths = []

    async def extract(path):
        paths.append(Path(path))
        return await original(path)

    monkeypatch.setattr(public, "extract_text", extract)
    result = await public.extract_chat_attachment(UploadFile(filename="policy.TXT", file=BytesIO(("正" * 16_001).encode())))
    assert len(result["content"]) == 16_000
    assert result["truncated"] is True
    assert result["type"] == ".txt"
    assert paths and all(not path.exists() and not path.parent.exists() for path in paths)


@pytest.mark.parametrize("payload", [{"request_id": "invalid id"}, {"request_id": "x" * 65}, {"files": [{"content": "x" * 16_001}]}, {"files": [{"content": "ok"}] * 6}])
def test_chat_query_bounds_identifiers_and_attachment_context(payload):
    with pytest.raises(ValidationError):
        ChatQuerySchema(message="hello", **payload)


@pytest.mark.parametrize("allowed", [True, False])
async def test_history_filters_revoked_citations_without_mutating_stored_runs(monkeypatch, allowed):
    citations = [{"id": "1", "title": "private.md", "snippet": "private knowledge", "knowledge_base_id": 7, "document_id": 9}]
    session = ChatSession(session_id="session-1", user_id="1", runs=[None, {"messages": [None, {"role": "assistant", "content": "answer", "citations": citations}]}])
    auth = SimpleNamespace(user=SimpleNamespace(id=1), db=object())

    async def accessible(current_auth, ids):
        assert current_auth is auth and ids == [7]
        if not allowed:
            raise CustomException(msg="知识库不存在或无权访问", status_code=403)
        return ids

    monkeypatch.setattr(service, "accessible_knowledge_base_ids", accessible)
    result = await service._format_session_data(session, auth)
    expected = citations if allowed else []
    assert result["messages"][0]["citations"] == expected
    assert result["runs"][1]["messages"][1]["citations"] == expected
    assert session.runs[1]["messages"][1]["citations"] == citations
    if not allowed:
        assert "private knowledge" not in json.dumps(result)


async def test_history_keeps_attachment_citations_without_requiring_a_knowledge_base(monkeypatch):
    citations = [{"id": "1", "title": "temporary.md", "snippet": "attachment text"}]
    session = ChatSession(session_id="session-1", user_id="1", runs=[{"messages": [{"role": "assistant", "content": "answer", "citations": citations}]}])
    accessible = AsyncMock(return_value=[])
    monkeypatch.setattr(service, "accessible_knowledge_base_ids", accessible)
    auth = SimpleNamespace(user=SimpleNamespace(id=1), db=object())
    result = await service._format_session_data(session, auth)
    assert result["messages"][0]["citations"] == citations
    accessible.assert_awaited_once_with(auth, [])
