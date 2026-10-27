import io
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import FastAPI, UploadFile
from httpx import ASGITransport, AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import event, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from app.api.v1.module_system.role.controller import RoleRouter
from app.api.v1.module_system.role.model import RoleModel
from app.api.v1.module_system.role.schema import RoleOutSchema
from app.api.v1.module_system.role.service import RoleService
from app.api.v1.module_system.user.controller import UserRouter, import_user_list_controller
from app.api.v1.module_system.user.model import UserModel
from app.api.v1.module_system.user.schema import UserOutSchema
from app.api.v1.module_system.user.service import UserService
from app.config.setting import settings
from app.core.base_model import MappedBase
from app.core.base_schema import AuthSchema, BatchSetAvailable
from app.core.dependencies import get_current_user
from app.core.exceptions import CustomException
from app.init_app import register_exceptions
from app.utils.hash_bcrpy_util import PwdUtil


def excel_upload(rows, template=None):
    workbook = load_workbook(io.BytesIO(template)) if template else Workbook()
    worksheet = workbook.active
    if not template:
        worksheet.append(["账号", "昵称", "邮箱", "手机号", "性别", "状态"])
    for row in rows:
        worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return UploadFile(filename="users.xlsx", file=buffer)


@pytest.fixture
async def batch_auth(tmp_path):
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'batch.db'}")
    async with engine.begin() as connection:
        await connection.run_sync(MappedBase.metadata.create_all)
    session_factory = async_sessionmaker(engine, expire_on_commit=False)
    async with session_factory() as session:
        yield AuthSchema(db=session, redis=AsyncMock(), check_data_scope=False), session_factory
    await engine.dispose()


async def test_batch_status_fetches_users_once_and_accepts_duplicate_ids(batch_auth):
    auth, session_factory = batch_auth
    users = [UserModel(username=f"batch{index}", name="用户", password="synthetic", status=0) for index in range(3)]
    auth.db.add_all(users)
    await auth.db.commit()
    ids = [user.id for user in users]
    statements = []

    def capture(_connection, _cursor, statement, _parameters, _context, _executemany):
        if statement.lstrip().upper().startswith("SELECT") and "FROM sys_user \nWHERE sys_user.is_deleted" in statement:
            statements.append(statement)

    engine = auth.db.bind.sync_engine
    event.listen(engine, "before_cursor_execute", capture)
    try:
        await UserService(auth).set_available(BatchSetAvailable(ids=[*ids, ids[0]], status=1))
        await auth.db.commit()
    finally:
        event.remove(engine, "before_cursor_execute", capture)

    async with session_factory() as persisted:
        assert list(await persisted.scalars(select(UserModel.status).order_by(UserModel.id))) == [1, 1, 1]
    assert len(statements) == 1, statements


@pytest.mark.parametrize(
    ("blocked", "message"),
    [
        ("missing", "该数据不存在"),
        ("deleted", "该数据不存在"),
        ("foreign", "该数据不存在"),
        ("superuser", "超级管理员状态不能修改"),
    ],
)
async def test_batch_status_validates_all_users_before_writing(batch_auth, blocked, message):
    auth, session_factory = batch_auth
    first = UserModel(username="first", name="正常用户", password="synthetic", created_id=1, status=0)
    second = UserModel(
        username="blocked",
        name="受保护用户",
        password="synthetic",
        created_id=2 if blocked == "foreign" else 1,
        is_deleted=blocked == "deleted",
        is_superuser=blocked == "superuser",
        status=0,
    )
    auth.db.add_all([first, second])
    await auth.db.commit()
    auth.user = SimpleNamespace(id=1, is_superuser=False, roles=[])
    auth.check_data_scope = True
    ids = [first.id, second.id + 1000 if blocked == "missing" else second.id]

    with pytest.raises(CustomException, match=message):
        await UserService(auth).set_available(BatchSetAvailable(ids=ids, status=1))
    await auth.db.commit()
    async with session_factory() as persisted:
        assert list(await persisted.scalars(select(UserModel.status).order_by(UserModel.id))) == [0, 0]


async def test_http_batch_status_updates_persist_and_protect_superadmin(batch_auth, monkeypatch):
    auth, session_factory = batch_auth
    auth.user = SimpleNamespace(id=None, is_superuser=True, roles=[])
    monkeypatch.setattr(settings, "OPERATION_LOG_RECORD", False)
    user = UserModel(username="httpbatch", name="接口用户", password="synthetic", status=0)
    admin = UserModel(username="protected", name="管理员", password="synthetic", is_superuser=True, status=0)
    auth.db.add_all([user, admin])
    await auth.db.commit()
    application = FastAPI()
    application.include_router(UserRouter, prefix="/system")
    application.dependency_overrides[get_current_user] = lambda: auth
    register_exceptions(application)
    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as client:
        response = await client.patch("/system/user/status/batch", json={"ids": [user.id, user.id], "status": 1})
        assert response.status_code == 200 and response.json()["success"] is True
        await auth.db.commit()
        response = await client.patch("/system/user/status/batch", json={"ids": [user.id, admin.id], "status": 0})
        assert response.status_code == 500 and response.json()["success"] is False
        assert response.json()["msg"] == "超级管理员状态不能修改"
        await auth.db.commit()
    async with session_factory() as persisted:
        assert list(await persisted.scalars(select(UserModel.status).order_by(UserModel.id))) == [1, 0]


@pytest.mark.parametrize("as_dict", [False, True])
def test_user_and_role_exports_are_real_excel_and_preserve_inputs(as_dict):
    users = [UserOutSchema(username=f"user{index}", name="用户", gender=str(index), status=0) for index in range(3)]
    role = RoleOutSchema(name="操作员", code="operator", data_scope=4, status=0)
    user_rows = [user.model_dump() for user in users] if as_dict else users
    role_rows = [role.model_dump()] if as_dict else [role]
    before_users = [dict(row) for row in user_rows] if as_dict else [user.model_dump() for user in users]
    before_role = dict(role_rows[0]) if as_dict else role.model_dump()

    user_workbook = load_workbook(io.BytesIO(UserService.export_list(user_rows)))
    user_values = list(user_workbook.active.values)
    gender_column = user_values[0].index("性别")
    assert [row[gender_column] for row in user_values[1:]] == ["男", "女", "未知"]
    assert all(row[user_values[0].index("状态")] == "启用" for row in user_values[1:])
    assert before_users == ([dict(row) for row in user_rows] if as_dict else [user.model_dump() for user in users])

    role_workbook = load_workbook(io.BytesIO(RoleService.export_list(role_rows)))
    role_values = list(role_workbook.active.values)
    assert role_values[1][role_values[0].index("数据权限")] == "全部数据权限"
    assert role_values[1][role_values[0].index("状态")] == "启用"
    assert before_role == (role_rows[0] if as_dict else role.model_dump())


async def test_downloaded_template_imports_gender_labels_and_persists(batch_auth):
    auth, session_factory = batch_auth
    upload = excel_upload(
        [[f"person{index}", "用户", None, None, gender, "正常"] for index, gender in enumerate(["男", "女", "未知"])],
        template=UserService.get_import_template(),
    )
    result = await UserService(auth).batch_import(upload)
    assert result["success_count"] == 3
    assert result["failed_count"] == 0
    assert result["errors"] == []
    assert result["message"] == "成功导入 3 条数据"
    await auth.db.commit()
    async with session_factory() as persisted:
        users = (await persisted.scalars(select(UserModel).order_by(UserModel.username))).all()
        assert [user.gender for user in users] == ["0", "1", "2"]


async def test_partial_import_isolates_constraint_failure_and_reports_excel_rows(batch_auth):
    auth, session_factory = batch_auth
    result = await UserService(auth).batch_import(excel_upload([
        ["first", "第一行", "same@example.com", None, 0, "正常"],
        ["duplicate", "重复邮箱", "same@example.com", None, 1, "正常"],
        ["last", "后续合法行", "last@example.com", None, 2, "正常"],
        ["invalid", "非法性别", None, None, "其他", "正常"],
        ["missing", None, None, None, "男", "正常"],
    ]))
    assert result["success_count"] == 2
    assert result["failed_count"] == 3
    assert [error["row"] for error in result["errors"]] == [3, 5, 6]
    assert "same@example.com" not in result["errors"][0]["message"]
    await auth.db.commit()
    async with session_factory() as persisted:
        assert list(await persisted.scalars(select(UserModel.username).order_by(UserModel.username))) == ["first", "last"]


async def test_import_updates_profile_without_changing_password_or_superadmin(batch_auth):
    auth, _ = batch_auth
    password_hash = PwdUtil.hash_password("Existing123")
    auth.db.add_all([
        UserModel(username="existing", name="原姓名", password=password_hash, gender="0"),
        UserModel(username="super", name="超级管理员", password=password_hash, is_superuser=True),
    ])
    await auth.db.commit()
    result = await UserService(auth).batch_import(excel_upload([
        ["existing", "新姓名", None, None, "女", "正常"],
        ["super", "不能修改", None, None, "男", "正常"],
    ]), update_support=True)
    assert (result["success_count"], result["failed_count"]) == (1, 1)
    assert result["errors"] == [{"row": 3, "message": "超级管理员不允许修改"}]
    existing = await auth.db.scalar(select(UserModel).where(UserModel.username == "existing"))
    assert existing.name == "新姓名"
    assert existing.gender == "1"
    assert existing.password == password_hash


@pytest.mark.parametrize("gender, expected", [(0, "0"), (1, "1"), (2, "2"), (None, "2")])
async def test_numeric_or_blank_gender_follows_schema(batch_auth, gender, expected):
    auth, _ = batch_auth
    result = await UserService(auth).batch_import(excel_upload([["numeric", "用户", None, None, gender, "正常"]]))
    assert result["success_count"] == 1
    assert (await auth.db.scalar(select(UserModel))).gender == expected


@pytest.mark.parametrize("valid_rows, failed_rows, expected_message", [(1, 0, "导入用户成功"), (1, 1, "部分导入失败"), (0, 1, "导入用户失败")])
async def test_import_controller_returns_counts_and_accurate_status(batch_auth, valid_rows, failed_rows, expected_message):
    auth, _ = batch_auth
    rows = [["valid", "合法", None, None, "男", "正常"]] * valid_rows
    rows += [["invalid", "非法", None, None, "其他", "正常"]] * failed_rows
    response = await import_user_list_controller(file=excel_upload(rows), auth=auth)
    payload = json.loads(response.body)
    assert response.status_code == 200
    assert payload["msg"] == expected_message
    assert payload["success"] is bool(valid_rows)
    assert (payload["data"]["success_count"], payload["data"]["failed_count"]) == (valid_rows, failed_rows)
    assert len(payload["data"]["errors"]) == failed_rows


async def test_http_upload_then_user_and_role_excel_downloads(batch_auth, monkeypatch):
    auth, _ = batch_auth
    auth.user = SimpleNamespace(id=None, is_superuser=True, roles=[])
    monkeypatch.setattr(settings, "OPERATION_LOG_RECORD", False)
    auth.db.add(RoleModel(name="导出角色", code="batch_role", data_scope=4, status=0))
    await auth.db.commit()
    application = FastAPI()
    application.include_router(UserRouter, prefix="/system")
    application.include_router(RoleRouter, prefix="/system")
    application.dependency_overrides[get_current_user] = lambda: auth
    upload = excel_upload([["httpuser", "接口用户", None, None, "男", "正常"]], template=UserService.get_import_template())
    async with AsyncClient(transport=ASGITransport(app=application), base_url="http://test") as client:
        response = await client.post("/system/user/import/data", files={"file": (upload.filename, upload.file, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")})
        assert response.status_code == 200
        assert response.json()["data"]["success_count"] == 1
        await auth.db.commit()
        for path, label, expected in [("user", "性别", "男"), ("role", "数据权限", "全部数据权限")]:
            response = await client.get(f"/system/{path}/export")
            assert response.status_code == 200
            assert "spreadsheetml.sheet" in response.headers["content-type"]
            values = list(load_workbook(io.BytesIO(response.content)).active.values)
            assert values[1][values[0].index(label)] == expected
